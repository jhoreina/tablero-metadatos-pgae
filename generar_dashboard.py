# -*- coding: utf-8 -*-
"""
generar_dashboard.py
--------------------
Script de automatización para generar o actualizar el dashboard de metadatos PGAE - SDA
en formato HTML/JS/CSS a partir de cualquier archivo Excel del modelo (soporta v10, vf, v9, etc.).

Uso:
    python generar_dashboard.py [ruta_al_excel.xlsx]
"""

import sys
import os
import json
import glob
import openpyxl

def parse_excel_metadata(excel_path):
    print(f"[+] Cargando archivo Excel: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. PGAE - Inventario (Atributos gobernados)
    sheet_inv_name = 'PGAE - Inventario' if 'PGAE - Inventario' in wb.sheetnames else [s for s in wb.sheetnames if 'Inventario' in s or 'pgae' in s.lower()][0]
    ws_inv = wb[sheet_inv_name]
    
    # Detectar dinámicamente las columnas por el nombre del encabezado
    header_row_idx = 3
    col_map = {}
    for c in range(1, ws_inv.max_column + 1):
        val = ws_inv.cell(header_row_idx, c).value
        if not val:
            continue
        v = str(val).lower().strip()
        if 'id' in v and 'atributo' in v: col_map['id'] = c
        elif 'eslab' in v: col_map['eslabon'] = c
        elif 'campo' in v: col_map['campo'] = c
        elif 'entidad' in v: col_map['entidad'] = c
        elif 'dato' in v or 'atributo base' in v: col_map['nombre'] = c
        elif 'naturaleza' in v: col_map['naturaleza'] = c
        elif 'grano' in v: col_map['grano'] = c
        elif 'frecuencia' in v: col_map['frecuencia'] = c
        elif 'fuente' in v and 'origin' in v: col_map['fuente'] = c
        elif 'sistema' in v: col_map['sistema'] = c
        elif 'área' in v or 'area' in v: col_map['area'] = c
        elif 'mecanismo' in v: col_map['acceso'] = c
        elif 'estado' in v: col_map['estado'] = c
        elif 'llave' in v or 'nivel' in v: col_map['llave'] = c
        elif 'replica' in v: col_map['replica'] = c
        elif 'grupo' in v: col_map['grupo'] = c
        elif 'pregunta' in v: col_map['preguntas'] = c

    # Mapeos por posición por defecto si no se detectan por nombre
    def get_cell(r, key, default_col):
        c = col_map.get(key, default_col)
        return str(ws_inv.cell(r, c).value or '').strip()

    attributes = []
    for r in range(4, ws_inv.max_row + 1):
        id_val = ws_inv.cell(r, col_map.get('id', 1)).value
        if id_val and str(id_val).strip().startswith('ATR-'):
            # Detectar si existe columna Entidad Maestra
            has_entidad = 'entidad' in col_map
            entidad_val = get_cell(r, 'entidad', 4) if has_entidad else 'N/A'
            
            # Ajuste de índices por defecto si existe o no Entidad Maestra
            offset = 1 if has_entidad else 0
            
            attributes.append({
                'id': str(id_val).strip(),
                'eslabon': get_cell(r, 'eslabon', 2),
                'campo': get_cell(r, 'campo', 3),
                'entidad': entidad_val,
                'nombre': get_cell(r, 'nombre', 4 + offset),
                'naturaleza': get_cell(r, 'naturaleza', 5 + offset),
                'grano': get_cell(r, 'grano', 6 + offset),
                'frecuencia': get_cell(r, 'frecuencia', 7 + offset),
                'fuente': get_cell(r, 'fuente', 8 + offset),
                'sistema': get_cell(r, 'sistema', 9 + offset),
                'area': get_cell(r, 'area', 10 + offset),
                'acceso': get_cell(r, 'acceso', 11 + offset),
                'estado': get_cell(r, 'estado', 12 + offset),
                'llave': get_cell(r, 'llave', 13 + offset),
                'replica': get_cell(r, 'replica', 14 + offset),
                'grupo': get_cell(r, 'grupo', 15 + offset),
                'preguntas': get_cell(r, 'preguntas', 16 + offset)
            })
    print(f"   OK -> Atributos extraidos: {len(attributes)} (Columna Entidad Maestra detectada: {'Sí' if 'entidad' in col_map else 'No'})")

    # 2. Matriz Eslabón x Campo
    ws_mat = wb['Matriz Eslabón x Campo']
    matrix = {}
    campos_header = ['CA-01', 'CA-02', 'CA-03', 'CA-04', 'CA-05', 'CA-06', 'CA-07', 'CA-08', 'CA-09']
    for r in range(5, 13):
        row_title = ws_mat.cell(r, 1).value
        if row_title and str(row_title).strip().startswith('E'):
            e_code = str(row_title).strip().split(' ')[0]
            matrix[e_code] = {}
            for idx, c_code in enumerate(campos_header, start=2):
                val = ws_mat.cell(r, idx).value or 0
                try:
                    matrix[e_code][c_code] = int(val)
                except ValueError:
                    matrix[e_code][c_code] = 0
    print("   OK -> Matriz Firma de Datos procesada")

    # 3. Campos de Acción
    ws_c = wb['Campos de Acción']
    campos = []
    for r in range(4, ws_c.max_row + 1):
        code = ws_c.cell(r, 1).value
        if code and str(code).strip().startswith('CA-'):
            c_code_clean = str(code).strip()
            count = sum(1 for a in attributes if c_code_clean in a['campo'])
            campos.append({
                'code': c_code_clean,
                'name': str(ws_c.cell(r, 2).value or '').strip(),
                'scope': str(ws_c.cell(r, 5).value or ws_c.cell(r, 4).value or '').strip(),
                'desc': str(ws_c.cell(r, 3).value or ws_c.cell(r, 4).value or '').strip(),
                'count': count,
                'icon': get_campo_icon(c_code_clean)
            })

    # 4. Eslabones Cadena de Valor
    ws_e = wb['Eslabones Cadena de Valor']
    eslabones = []
    for r in range(4, ws_e.max_row + 1):
        code = ws_e.cell(r, 1).value
        if code and str(code).strip().startswith('E'):
            e_code_clean = str(code).strip()
            count = sum(1 for a in attributes if e_code_clean in a['eslabon'])
            eslabones.append({
                'code': e_code_clean,
                'name': str(ws_e.cell(r, 2).value or '').strip(),
                'deliverable': str(ws_e.cell(r, 5).value or ws_e.cell(r, 4).value or '').strip(),
                'count': count,
                'icon': get_eslabon_icon(e_code_clean)
            })

    # 5. Preguntas Estratégicas
    ws_p = wb['Preguntas Estratégicas']
    preguntas = []
    for r in range(3, ws_p.max_row + 1):
        code = ws_p.cell(r, 1).value
        if code and str(code).strip().startswith('PE-'):
            preguntas.append({
                'code': str(code).strip(),
                'topic': str(ws_p.cell(r, 2).value or '').strip(),
                'question': str(ws_p.cell(r, 3).value or '').strip(),
                'ca': str(ws_p.cell(r, 4).value or '').strip()
            })

    # 6. Dimensiones Maestras
    ws_d = wb['Dimensiones Maestras']
    dimensiones = []
    for r in range(3, ws_d.max_row + 1):
        code = ws_d.cell(r, 1).value
        if code and str(code).strip().startswith('DM-'):
            dimensiones.append({
                'id': str(code).strip(),
                'name': str(ws_d.cell(r, 2).value or '').strip(),
                'concept': str(ws_d.cell(r, 3).value or '').strip(),
                'attributes': str(ws_d.cell(r, 4).value or '').strip(),
                'system': str(ws_d.cell(r, 5).value or '').strip(),
                'owner': str(ws_d.cell(r, 6).value or '').strip(),
                'uses': str(ws_d.cell(r, 7).value or '').strip()
            })

    # 7. Indicadores Derivados
    ws_ind = wb['Indicadores Derivados']
    indicadores = []
    for r in range(4, ws_ind.max_row + 1):
        name = ws_ind.cell(r, 1).value
        if name:
            indicadores.append({
                'name': str(name).strip(),
                'ca': str(ws_ind.cell(r, 2).value or '').strip(),
                'formula': str(ws_ind.cell(r, 3).value or '').strip(),
                'inputs': str(ws_ind.cell(r, 4).value or '').strip(),
                'freq': str(ws_ind.cell(r, 5).value or '').strip(),
                'role': str(ws_ind.cell(r, 6).value or '').strip(),
                'use': str(ws_ind.cell(r, 7).value or '').strip()
            })

    # 8. Catálogo de Fuentes
    ws_f = wb['Catálogo de Fuentes']
    fuentes = []
    for r in range(4, ws_f.max_row + 1):
        num = ws_f.cell(r, 1).value
        if num is not None:
            fuentes.append({
                'id': str(num).strip(),
                'name': str(ws_f.cell(r, 2).value or '').strip(),
                'type': str(ws_f.cell(r, 3).value or '').strip(),
                'attrs': str(ws_f.cell(r, 4).value or '').strip(),
                'owner': str(ws_f.cell(r, 5).value or '').strip(),
                'mechanism': str(ws_f.cell(r, 6).value or '').strip(),
                'group': str(ws_f.cell(r, 7).value or '').strip(),
                'status': str(ws_f.cell(r, 8).value or '').strip(),
                'lead': str(ws_f.cell(r, 9).value or '').strip()
            })

    # 9. Roadmap
    roadmap = [
        {'phase': 'Fase 1: Consolidación Máster', 'horizon': 'Mes 1', 'deliverable': 'Despliegue de dimensiones maestras (Empresa y Sede con clave GAE_NIT_Consecutivo) y migración de inventarios GAE-PREAD.', 'systems': 'Ventanilla Virtual, GAE PostgreSQL', 'success': 'Cero sedes huérfanas y asignación de identificador único'},
        {'phase': 'Fase 2: Interoperabilidad Distrital', 'horizon': 'Mes 2', 'deliverable': 'Conexión web service con BOG Data (SDH) para validación de CHIPs prediales y con SINUPOT para conceptos de suelo.', 'systems': 'BOG Data, SINUPOT, Forest PR46', 'success': 'Validación automática de predios e informes PR46 adjuntos'},
        {'phase': 'Fase 3: Interoperabilidad Nacional y ESP', 'horizon': 'Mes 3', 'deliverable': 'Integración batch con RUA del IDEAM y convenios de datos directos con empresas de servicios públicos (EAAB, Enel, Vanti).', 'systems': 'RUA IDEAM, Sistemas ESP', 'success': 'Sincronización de balances de agua/energía sin digitación manual'}
    ]

    return {
        'attributes': attributes,
        'matrix': matrix,
        'campos': campos,
        'eslabones': eslabones,
        'preguntas': preguntas,
        'dimensiones': dimensiones,
        'indicadores': indicadores,
        'fuentes': fuentes,
        'roadmap': roadmap
    }

def get_campo_icon(code):
    icons = {
        'CA-01': '🏢', 'CA-02': '🏭', 'CA-03': '💧', 'CA-04': '⚡',
        'CA-05': '☣️', 'CA-06': '🌱', 'CA-07': '💼', 'CA-08': '⚖️', 'CA-09': '📄'
    }
    return icons.get(code, '📌')

def get_eslabon_icon(code):
    icons = {
        'E1': '📝', 'E2': '🗺️', 'E3': '📊', 'E4': '🤝',
        'E5': '🔍', 'E6': '📈', 'E7': '🏆', 'E8': '🔄'
    }
    return icons.get(code, '🔹')

def generate_html_dashboard(data, output_path):
    print(f"[+] Compilando portal HTML en: {output_path}...")
    
    attributes_json = json.dumps(data['attributes'], ensure_ascii=False)
    campos_json = json.dumps(data['campos'], ensure_ascii=False)
    eslabones_json = json.dumps(data['eslabones'], ensure_ascii=False)
    preguntas_json = json.dumps(data['preguntas'], ensure_ascii=False)
    dimensiones_json = json.dumps(data['dimensiones'], ensure_ascii=False)
    indicadores_json = json.dumps(data['indicadores'], ensure_ascii=False)
    fuentes_json = json.dumps(data['fuentes'], ensure_ascii=False)
    matrix_json = json.dumps(data['matrix'], ensure_ascii=False)
    roadmap_json = json.dumps(data['roadmap'], ensure_ascii=False)
    
    total_attrs = len(data['attributes'])

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Modelo de Metadatos y Gobernanza de Datos PGAE - SDA</title>
    <!-- Chart.js para visualización de métricas -->
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <!-- SheetJS para carga de archivos Excel en vivo desde el navegador -->
    <script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
    <style>
        :root {{
            --sda-dark-green: #577515;
            --sda-bright-green: #89BC24;
            --sda-bg-light: #F9F9F9;
            --sda-card-bg: #FFFFFF;
            --sda-text-main: #2D3748;
            --sda-text-muted: #64748B;
            --sda-border: #E2E8F0;
            --sda-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03);
            --sda-radius: 10px;
        }}

        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
        }}

        body {{
            font-family: 'Calibri', 'Segoe UI', Arial, sans-serif;
            background-color: var(--sda-bg-light);
            color: var(--sda-text-main);
            line-height: 1.5;
        }}

        /* Header Institucional */
        header {{
            background: linear-gradient(135deg, var(--sda-dark-green) 0%, #3e550e 100%);
            color: white;
            padding: 1.5rem 2rem;
            border-bottom: 4px solid var(--sda-bright-green);
            box-shadow: 0 4px 12px rgba(0,0,0,0.15);
        }}

        .header-container {{
            max-width: 1400px;
            margin: 0 auto;
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 1rem;
        }}

        .header-title h1 {{
            font-size: 1.8rem;
            font-weight: 700;
            letter-spacing: 0.5px;
        }}

        .header-title p {{
            font-size: 1.05rem;
            opacity: 0.9;
            color: #dbe8be;
        }}

        .header-actions {{
            display: flex;
            gap: 12px;
            align-items: center;
        }}

        .header-badge {{
            background: var(--sda-bright-green);
            color: #1a3000;
            padding: 0.5rem 1.2rem;
            border-radius: 20px;
            font-weight: bold;
            font-size: 0.9rem;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}

        .btn-upload {{
            background: #ffffff;
            color: var(--sda-dark-green);
            padding: 0.5rem 1.1rem;
            border-radius: 20px;
            font-weight: bold;
            font-size: 0.85rem;
            cursor: pointer;
            border: 2px solid var(--sda-bright-green);
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}

        .btn-upload:hover {{
            background: #f4f8ec;
            transform: translateY(-1px);
        }}

        /* Bar Nav Tab */
        nav {{
            background-color: white;
            border-bottom: 1px solid var(--sda-border);
            position: sticky;
            top: 0;
            z-index: 100;
            box-shadow: 0 2px 4px rgba(0,0,0,0.04);
        }}

        .nav-container {{
            max-width: 1400px;
            margin: 0 auto;
            display: flex;
            overflow-x: auto;
            white-space: nowrap;
            scrollbar-width: thin;
        }}

        .nav-tab {{
            padding: 1rem 1.25rem;
            font-size: 0.95rem;
            font-weight: 600;
            color: var(--sda-text-muted);
            border: none;
            background: none;
            cursor: pointer;
            border-bottom: 3px solid transparent;
            transition: all 0.2s ease;
            display: flex;
            align-items: center;
            gap: 8px;
        }}

        .nav-tab:hover {{
            color: var(--sda-dark-green);
            background-color: #f4f8ec;
        }}

        .nav-tab.active {{
            color: var(--sda-dark-green);
            border-bottom-color: var(--sda-bright-green);
            background-color: #f4f8ec;
        }}

        /* Main Body */
        main {{
            max-width: 1400px;
            margin: 2rem auto;
            padding: 0 1.5rem;
        }}

        .tab-content {{
            display: none;
        }}

        .tab-content.active {{
            display: block;
            animation: fadeIn 0.3s ease-in-out;
        }}

        @keyframes fadeIn {{
            from {{ opacity: 0; transform: translateY(6px); }}
            to {{ opacity: 1; transform: translateY(0); }}
        }}

        /* Grid Metrics Cards */
        .kpi-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 1.25rem;
            margin-bottom: 2rem;
        }}

        .kpi-card {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.25rem;
            box-shadow: var(--sda-shadow);
            border-top: 4px solid var(--sda-bright-green);
            display: flex;
            flex-direction: column;
            justify-content: space-between;
        }}

        .kpi-card.accent {{
            border-top-color: var(--sda-dark-green);
        }}

        .kpi-title {{
            font-size: 0.85rem;
            text-transform: uppercase;
            letter-spacing: 0.5px;
            color: var(--sda-text-muted);
            font-weight: 700;
        }}

        .kpi-value {{
            font-size: 2.2rem;
            font-weight: 800;
            color: var(--sda-dark-green);
            margin: 0.4rem 0;
        }}

        .kpi-sub {{
            font-size: 0.8rem;
            color: var(--sda-text-muted);
        }}

        /* Section Title */
        .section-header {{
            margin-bottom: 1.5rem;
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 1rem;
        }}

        .section-title {{
            font-size: 1.4rem;
            color: var(--sda-dark-green);
            font-weight: 700;
            display: flex;
            align-items: center;
            gap: 10px;
        }}

        /* Matrix Table */
        .matrix-container {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.5rem;
            box-shadow: var(--sda-shadow);
            margin-bottom: 2rem;
            overflow-x: auto;
        }}

        .matrix-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 0.9rem;
            text-align: center;
        }}

        .matrix-table th, .matrix-table td {{
            padding: 0.75rem 0.5rem;
            border: 1px solid var(--sda-border);
        }}

        .matrix-table th {{
            background-color: #f2f7e9;
            color: var(--sda-dark-green);
            font-weight: 700;
        }}

        .matrix-cell {{
            font-weight: bold;
            cursor: pointer;
            transition: transform 0.15s ease, box-shadow 0.15s ease;
        }}

        .matrix-cell:hover {{
            transform: scale(1.08);
            box-shadow: 0 4px 8px rgba(0,0,0,0.15);
            z-index: 10;
        }}

        .cell-0 {{ background-color: #f8fafc; color: #94a3b8; }}
        .cell-low {{ background-color: #e2f2cc; color: #3b520a; }}
        .cell-med {{ background-color: #c1e685; color: #283a04; }}
        .cell-high {{ background-color: #89bc24; color: #152400; }}
        .cell-dark {{ background-color: #577515; color: #ffffff; }}

        /* Charts Layout */
        .charts-row {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(400px, 1fr));
            gap: 1.5rem;
            margin-bottom: 2rem;
        }}

        .chart-card {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.5rem;
            box-shadow: var(--sda-shadow);
        }}

        .chart-card h3 {{
            font-size: 1.1rem;
            color: var(--sda-dark-green);
            margin-bottom: 1rem;
        }}

        /* Timeline Roadmap */
        .roadmap-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(300px, 1fr));
            gap: 1.25rem;
            margin-bottom: 2rem;
        }}

        .roadmap-card {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.25rem;
            box-shadow: var(--sda-shadow);
            border-left: 5px solid var(--sda-bright-green);
        }}

        .roadmap-phase {{
            font-size: 0.8rem;
            font-weight: 700;
            text-transform: uppercase;
            color: var(--sda-bright-green);
        }}

        .roadmap-title {{
            font-size: 1.1rem;
            font-weight: 700;
            color: var(--sda-dark-green);
            margin: 0.3rem 0 0.5rem 0;
        }}

        .roadmap-desc {{
            font-size: 0.9rem;
            color: var(--sda-text-main);
            margin-bottom: 0.8rem;
        }}

        .roadmap-meta {{
            font-size: 0.8rem;
            background: #f1f5f9;
            padding: 0.5rem;
            border-radius: 6px;
            color: var(--sda-text-muted);
        }}

        /* Controls */
        .controls-card {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.25rem;
            box-shadow: var(--sda-shadow);
            margin-bottom: 1.5rem;
            display: flex;
            flex-wrap: wrap;
            gap: 1rem;
            align-items: center;
            justify-content: space-between;
        }}

        .search-box {{
            flex: 1;
            min-width: 260px;
            position: relative;
        }}

        .search-box input {{
            width: 100%;
            padding: 0.65rem 1rem 0.65rem 2.4rem;
            border: 1px solid var(--sda-border);
            border-radius: 6px;
            font-size: 0.95rem;
            font-family: inherit;
        }}

        .search-box::before {{
            content: "🔍";
            position: absolute;
            left: 0.8rem;
            top: 50%;
            transform: translateY(-50%);
            font-size: 0.9rem;
            opacity: 0.5;
        }}

        .filter-group {{
            display: flex;
            gap: 0.8rem;
            flex-wrap: wrap;
        }}

        .filter-select {{
            padding: 0.65rem 1rem;
            border: 1px solid var(--sda-border);
            border-radius: 6px;
            font-size: 0.9rem;
            font-family: inherit;
            background-color: white;
            color: var(--sda-text-main);
        }}

        .btn {{
            padding: 0.65rem 1.2rem;
            border-radius: 6px;
            font-weight: 600;
            font-size: 0.9rem;
            cursor: pointer;
            border: none;
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}

        .btn-primary {{
            background-color: var(--sda-dark-green);
            color: white;
        }}

        .btn-primary:hover {{
            background-color: #445d0f;
        }}

        .btn-outline {{
            background-color: transparent;
            border: 1px solid var(--sda-dark-green);
            color: var(--sda-dark-green);
        }}

        .btn-outline:hover {{
            background-color: #f4f8ec;
        }}

        /* Data Table */
        .table-card {{
            background: white;
            border-radius: var(--sda-radius);
            box-shadow: var(--sda-shadow);
            overflow: hidden;
            margin-bottom: 2rem;
        }}

        .data-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 0.9rem;
        }}

        .data-table th {{
            background-color: #f2f7e9;
            color: var(--sda-dark-green);
            padding: 0.85rem 1rem;
            text-align: left;
            font-weight: 700;
            border-bottom: 2px solid var(--sda-border);
        }}

        .data-table td {{
            padding: 0.85rem 1rem;
            border-bottom: 1px solid var(--sda-border);
            color: var(--sda-text-main);
        }}

        .data-table tbody tr:hover {{
            background-color: #f8faf5;
            cursor: pointer;
        }}

        .badge {{
            padding: 3px 8px;
            border-radius: 12px;
            font-size: 0.75rem;
            font-weight: bold;
            display: inline-block;
        }}

        .badge-original {{ background: #e0f2fe; color: #0369a1; }}
        .badge-homologado {{ background: #fef3c7; color: #b45309; }}
        .badge-validado {{ background: #dcfce7; color: #15803d; }}
        .badge-entidad {{ background: #f3e8ff; color: #6b21a8; font-weight: 600; }}

        /* Generic Cards Grid */
        .card-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
            gap: 1.25rem;
            margin-bottom: 2rem;
        }}

        .info-card {{
            background: white;
            border-radius: var(--sda-radius);
            padding: 1.5rem;
            box-shadow: var(--sda-shadow);
            border: 1px solid var(--sda-border);
            transition: transform 0.2s ease, box-shadow 0.2s ease;
        }}

        .info-card:hover {{
            transform: translateY(-3px);
            box-shadow: 0 8px 15px rgba(0,0,0,0.08);
        }}

        .info-header {{
            display: flex;
            align-items: center;
            gap: 12px;
            margin-bottom: 1rem;
        }}

        .info-icon {{
            font-size: 1.8rem;
            background: #f4f8ec;
            padding: 8px;
            border-radius: 8px;
            color: var(--sda-dark-green);
        }}

        .info-title {{
            font-size: 1.1rem;
            font-weight: 700;
            color: var(--sda-dark-green);
        }}

        .info-sub {{
            font-size: 0.8rem;
            color: var(--sda-bright-green);
            font-weight: bold;
            text-transform: uppercase;
        }}

        .info-body {{
            font-size: 0.9rem;
            color: var(--sda-text-main);
        }}

        .formula-box {{
            background: #f8fafc;
            border-left: 4px solid var(--sda-dark-green);
            padding: 0.75rem 1rem;
            font-family: 'Courier New', Courier, monospace;
            font-size: 0.9rem;
            margin: 0.75rem 0;
            border-radius: 0 6px 6px 0;
            color: #1e293b;
            font-weight: bold;
        }}

        /* Modal Styles */
        .modal-overlay {{
            position: fixed;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            background: rgba(0,0,0,0.5);
            display: none;
            justify-content: center;
            align-items: center;
            z-index: 1000;
            padding: 1rem;
        }}

        .modal-content {{
            background: white;
            border-radius: var(--sda-radius);
            max-width: 700px;
            width: 100%;
            max-height: 90vh;
            overflow-y: auto;
            box-shadow: 0 20px 25px -5px rgba(0, 0, 0, 0.1);
            animation: modalPop 0.25s ease-out;
        }}

        @keyframes modalPop {{
            from {{ transform: scale(0.95); opacity: 0; }}
            to {{ transform: scale(1); opacity: 1; }}
        }}

        .modal-header {{
            background: var(--sda-dark-green);
            color: white;
            padding: 1.25rem 1.5rem;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }}

        .modal-header h3 {{
            font-size: 1.2rem;
        }}

        .modal-close {{
            background: none;
            border: none;
            color: white;
            font-size: 1.5rem;
            cursor: pointer;
        }}

        .modal-body {{
            padding: 1.5rem;
        }}

        .detail-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
            gap: 1rem;
        }}

        .detail-item {{
            background: #f8fafc;
            padding: 0.85rem;
            border-radius: 6px;
            border: 1px solid var(--sda-border);
        }}

        .detail-label {{
            font-size: 0.75rem;
            text-transform: uppercase;
            color: var(--sda-text-muted);
            font-weight: 700;
        }}

        .detail-value {{
            font-size: 0.95rem;
            font-weight: 600;
            color: var(--sda-text-main);
            margin-top: 2px;
        }}

        /* Footer */
        footer {{
            background: #1e293b;
            color: #94a3b8;
            text-align: center;
            padding: 2rem 1rem;
            margin-top: 4rem;
            font-size: 0.85rem;
            border-top: 4px solid var(--sda-bright-green);
        }}

        footer a {{
            color: var(--sda-bright-green);
            text-decoration: none;
        }}

        @media (max-width: 768px) {{
            .header-container {{
                flex-direction: column;
                align-items: flex-start;
            }}
            .charts-row {{
                grid-template-columns: 1fr;
            }}
        }}
    </style>
</head>
<body>

    <!-- Header Institucional -->
    <header>
        <div class="header-container">
            <div class="header-title">
                <h1>Modelo de Datos y Metadatos PGAE (v1.0)</h1>
                <p>Programa de Gestión Ambiental Empresarial (PGAE)</p>
            </div>
            <div class="header-actions">
                <label class="btn-upload" title="Cargar un nuevo archivo Excel para actualizar la plataforma en tiempo real">
                    📁 Cargar Excel (.xlsx)
                    <input type="file" id="excelFileInput" accept=".xlsx" style="display:none;" onchange="handleExcelUpload(event)">
                </label>
                <div class="header-badge">
                    🌿 Subdirección SEGAE
                </div>
            </div>
        </div>
    </header>

    <!-- Barra de Navegación por Pestañas -->
    <nav>
        <div class="nav-container">
            <button class="nav-tab active" onclick="switchTab('tab-general')">📊 Visión General</button>
            <button class="nav-tab" onclick="switchTab('tab-inventario')">🗂️ Inventario de Atributos</button>
            <button class="nav-tab" onclick="switchTab('tab-campos')">🎯 Campos y Eslabones</button>
            <button class="nav-tab" onclick="switchTab('tab-preguntas')">❓ Preguntas Estratégicas</button>
            <button class="nav-tab" onclick="switchTab('tab-dimensiones')">🧩 Dimensiones Maestras</button>
            <button class="nav-tab" onclick="switchTab('tab-indicadores')">📈 Indicadores Derivados</button>
            <button class="nav-tab" onclick="switchTab('tab-fuentes')">💻 Catálogo de Fuentes</button>
        </div>
    </nav>

    <!-- Contenido Principal -->
    <main>

        <!-- PESTAÑA 1: VISIÓN GENERAL -->
        <div id="tab-general" class="tab-content active">
            <!-- Métricas KPI -->
            <div class="kpi-grid">
                <div class="kpi-card">
                    <div class="kpi-title">Total Atributos Gobernados</div>
                    <div class="kpi-value" id="kpiTotalAttrs">{total_attrs}</div>
                    <div class="kpi-sub">Inventario completo PGAE (3NF)</div>
                </div>
                <div class="kpi-card accent">
                    <div class="kpi-title">Campos de Acción</div>
                    <div class="kpi-value" id="kpiCampos">9</div>
                    <div class="kpi-sub">Vectores temáticos de empresa</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-title">Eslabones Cadena de Valor</div>
                    <div class="kpi-value" id="kpiEslabones">8</div>
                    <div class="kpi-sub">Etapas del proceso institucional</div>
                </div>
                <div class="kpi-card accent">
                    <div class="kpi-title">Preguntas Estratégicas</div>
                    <div class="kpi-value" id="kpiPreguntas">15</div>
                    <div class="kpi-sub">Requerimientos de analítica</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-title">Dimensiones Maestras</div>
                    <div class="kpi-value" id="kpiDimensiones">7</div>
                    <div class="kpi-sub">Claves únicas de gobernanza</div>
                </div>
                <div class="kpi-card accent">
                    <div class="kpi-title">Sistemas Fuentes</div>
                    <div class="kpi-value" id="kpiFuentes">9</div>
                    <div class="kpi-sub">Catálogo normalizado OTI</div>
                </div>
            </div>

            <!-- Matriz Firma de Datos -->
            <div class="section-header">
                <h2 class="section-title"> Matriz Eslabón × Campo de Acción ("Firma de Datos")</h2>
                <span style="font-size:0.85rem; color:var(--sda-text-muted);">Haz clic en cualquier celda para filtrar el inventario</span>
            </div>

            <div class="matrix-container">
                <table class="matrix-table" id="matrixTable">
                    <thead>
                        <tr>
                            <th>Eslabón de la Cadena</th>
                            <th>CA-01<br>Predial</th>
                            <th>CA-02<br>Productiva</th>
                            <th>CA-03<br>Hídrico</th>
                            <th>CA-04<br>Energía</th>
                            <th>CA-05<br>RESPEL</th>
                            <th>CA-06<br>SGA/DGA</th>
                            <th>CA-07<br>Proyectos</th>
                            <th>CA-08<br>Legalidad</th>
                            <th>CA-09<br>Soportes</th>
                            <th>Total</th>
                        </tr>
                    </thead>
                    <tbody id="matrixBody">
                        <!-- Generado por JavaScript -->
                    </tbody>
                </table>
            </div>

            <!-- Gráficos -->
            <div class="charts-row">
                <div class="chart-card">
                    <h3>Distribución de Atributos por Campo de Acción</h3>
                    <canvas id="chartCampos" height="220"></canvas>
                </div>
                <div class="chart-card">
                    <h3>Densidad de Atributos por Eslabón de Proceso</h3>
                    <canvas id="chartEslabones" height="220"></canvas>
                </div>
            </div>

            <!-- Hoja de Ruta OTI -->
            <div class="section-header">
                <h2 class="section-title">🗺️ Hoja de Ruta e Implementación Técnica (OTI)</h2>
            </div>
            <div class="roadmap-grid" id="roadmapGrid">
                <!-- Generado por JavaScript -->
            </div>
        </div>

        <!-- PESTAÑA 2: INVENTARIO DE ATRIBUTOS -->
        <div id="tab-inventario" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">🗂️ Catálogo e Inventario Gobernadode Atributos</h2>
                <div>
                    <button class="btn btn-outline" onclick="exportData('json')">📥 Exportar JSON</button>
                    <button class="btn btn-primary" onclick="exportData('csv')">📊 Exportar CSV</button>
                </div>
            </div>

            <!-- Filtros y Búsqueda -->
            <div class="controls-card">
                <div class="search-box">
                    <input type="text" id="searchInput" placeholder="Buscar por ID, atributo, entidad maestra, fuente, sistema..." onkeyup="filterInventory()">
                </div>
                <div class="filter-group">
                    <select id="filterCampo" class="filter-select" onchange="filterInventory()">
                        <option value="">Todos los Campos de Acción</option>
                    </select>
                    <select id="filterEslabon" class="filter-select" onchange="filterInventory()">
                        <option value="">Todos los Eslabones</option>
                    </select>
                    <select id="filterEntidad" class="filter-select" onchange="filterInventory()">
                        <option value="">Todas las Entidades Maestras</option>
                    </select>
                    <select id="filterEstado" class="filter-select" onchange="filterInventory()">
                        <option value="">Todos los Estados del Dato</option>
                        <option value="DATO ORIGINAL">Dato Original</option>
                        <option value="DATO HOMOLOGADO">Dato Homologado</option>
                        <option value="DATO VALIDADO">Dato Validado</option>
                    </select>
                </div>
            </div>

            <!-- Tabla de Datos -->
            <div class="table-card">
                <table class="data-table">
                    <thead>
                        <tr>
                            <th>ID Atributo</th>
                            <th>Atributo Base</th>
                            <th>Entidad Maestra</th>
                            <th>Campo de Acción</th>
                            <th>Eslabón Cadena</th>
                            <th>Sistema Fuente</th>
                            <th>Estado</th>
                        </tr>
                    </thead>
                    <tbody id="inventoryTableBody">
                        <!-- Generado por JavaScript -->
                    </tbody>
                </table>
            </div>
        </div>

        <!-- PESTAÑA 3: CAMPOS Y ESLABONES -->
        <div id="tab-campos" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">🎯 Campos de Acción (Vectores Temáticos)</h2>
            </div>
            <div class="card-grid" id="camposGrid"></div>

            <div class="section-header" style="margin-top: 2rem;">
                <h2 class="section-title">🔄 Eslabones de la Cadena de Valor (Procesos SDA)</h2>
            </div>
            <div class="card-grid" id="eslabonesGrid"></div>
        </div>

        <!-- PESTAÑA 4: PREGUNTAS ESTRATÉGICAS -->
        <div id="tab-preguntas" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">❓ Preguntas Estratégicas y Demanda de Información</h2>
            </div>
            <div class="card-grid" id="preguntasGrid"></div>
        </div>

        <!-- PESTAÑA 5: DIMENSIONES MAESTRAS -->
        <div id="tab-dimensiones" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">🧩 Dimensiones Maestras y Claves de Gobernanza</h2>
            </div>
            <div class="card-grid" id="dimensionesGrid"></div>
        </div>

        <!-- PESTAÑA 6: INDICADORES DERIVADOS -->
        <div id="tab-indicadores" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">📈 Indicadores Derivados y Reglas de Cómputo</h2>
            </div>
            <div class="card-grid" id="indicadoresGrid"></div>
        </div>

        <!-- PESTAÑA 7: CATÁLOGO DE FUENTES -->
        <div id="tab-fuentes" class="tab-content">
            <div class="section-header">
                <h2 class="section-title">💻 Catálogo de Fuentes Originarias y Accesos OTI</h2>
            </div>
            <div class="card-grid" id="fuentesGrid"></div>
        </div>

    </main>

    <!-- Modal Detalle Atributo -->
    <div class="modal-overlay" id="detailModal">
        <div class="modal-content">
            <div class="modal-header">
                <h3 id="modalTitle">Ficha del Atributo</h3>
                <button class="modal-close" onclick="closeModal()">&times;</button>
            </div>
            <div class="modal-body" id="modalBody">
                <!-- Generado dinámicamente -->
            </div>
        </div>
    </div>

    <!-- Footer -->
    <footer>
        <p>Alcaldía Mayor de Bogotá D.C. — Secretaría Distrital de Ambiente (SDA)</p>
        <p>Modelo de Datos y Metadatos PGAE v1.0 | Subdirección SEGAE | 2026</p>
    </footer>

    <!-- Script principal de datos e interacción -->
    <script>
        // Datos iniciales embebidos desde el Excel
        let attributesData = {attributes_json};
        let camposData = {campos_json};
        let eslabonesData = {eslabones_json};
        let preguntasData = {preguntas_json};
        let dimensionesData = {dimensiones_json};
        let indicadoresData = {indicadores_json};
        let fuentesData = {fuentes_json};
        let matrixData = {matrix_json};
        let roadmapData = {roadmap_json};

        let chartCamposInstance = null;
        let chartEslabonesInstance = null;

        // Inicialización al cargar la página
        document.addEventListener('DOMContentLoaded', () => {{
            refreshAllUI();
        }});

        function refreshAllUI() {{
            document.getElementById('kpiTotalAttrs').textContent = attributesData.length;
            document.getElementById('kpiCampos').textContent = camposData.length;
            document.getElementById('kpiEslabones').textContent = eslabonesData.length;
            document.getElementById('kpiPreguntas').textContent = preguntasData.length;
            document.getElementById('kpiDimensiones').textContent = dimensionesData.length;
            document.getElementById('kpiFuentes').textContent = fuentesData.length;

            renderMatrix();
            renderRoadmap();
            renderCampos();
            renderEslabones();
            renderPreguntas();
            renderDimensiones();
            renderIndicadores();
            renderFuentes();
            populateFilters();
            renderInventory(attributesData);
            initCharts();
        }}

        // Cambiar pestañas
        function switchTab(tabId) {{
            document.querySelectorAll('.tab-content').forEach(el => el.classList.remove('active'));
            document.querySelectorAll('.nav-tab').forEach(el => el.classList.remove('active'));
            
            document.getElementById(tabId).classList.add('active');
            event.currentTarget.classList.add('active');
        }}

        // Render Matriz Firma de Datos
        function renderMatrix() {{
            const tbody = document.getElementById('matrixBody');
            tbody.innerHTML = '';
            
            const campos = ['CA-01', 'CA-02', 'CA-03', 'CA-04', 'CA-05', 'CA-06', 'CA-07', 'CA-08', 'CA-09'];
            
            Object.keys(matrixData).forEach(eKey => {{
                let tr = document.createElement('tr');
                let eName = eslabonesData.find(e => e.code === eKey)?.name || eKey;
                let totalRow = 0;
                
                let html = `<td style="text-align:left; font-weight:bold;">${{eKey}} - ${{eName}}</td>`;
                
                campos.forEach(cKey => {{
                    let val = matrixData[eKey][cKey] || 0;
                    totalRow += val;
                    let cellClass = 'cell-0';
                    if (val >= 7) cellClass = 'cell-dark';
                    else if (val >= 4) cellClass = 'cell-high';
                    else if (val >= 2) cellClass = 'cell-med';
                    else if (val >= 1) cellClass = 'cell-low';
                    
                    html += `<td class="matrix-cell ${{cellClass}}" onclick="filterMatrixCell('${{eKey}}', '${{cKey}}')">${{val}}</td>`;
                }});
                
                html += `<td style="font-weight:bold; background:#e2e8f0;">${{totalRow}}</td>`;
                tr.innerHTML = html;
                tbody.appendChild(tr);
            }});

            // Row Totals
            let trTotals = document.createElement('tr');
            let totalsHtml = '<td style="font-weight:bold; background:#cbd5e1; text-align:left;">Total Atributos por Campo</td>';
            let grandTotal = 0;
            
            campos.forEach(cKey => {{
                let cTotal = 0;
                Object.keys(matrixData).forEach(eKey => {{ cTotal += matrixData[eKey][cKey] || 0; }});
                grandTotal += cTotal;
                totalsHtml += `<td style="font-weight:bold; background:#cbd5e1;">${{cTotal}}</td>`;
            }});
            totalsHtml += `<td style="font-weight:bold; background:#94a3b8; color:white;">${{grandTotal}}</td>`;
            trTotals.innerHTML = totalsHtml;
            tbody.appendChild(trTotals);
        }}

        // Filtrar por celda de matriz
        function filterMatrixCell(eCode, cCode) {{
            switchTab('tab-inventario');
            document.getElementById('filterEslabon').value = eslabonesData.find(e => e.code === eCode)?.name || '';
            document.getElementById('filterCampo').value = camposData.find(c => c.code === cCode)?.name || '';
            filterInventory();
        }}

        // Render Hoja de Ruta OTI
        function renderRoadmap() {{
            const grid = document.getElementById('roadmapGrid');
            grid.innerHTML = roadmapData.map(r => `
                <div class="roadmap-card">
                    <div class="roadmap-phase">${{r.horizon}}</div>
                    <div class="roadmap-title">${{r.phase}}</div>
                    <div class="roadmap-desc">${{r.deliverable}}</div>
                    <div class="roadmap-meta">
                        <strong>Sistemas:</strong> ${{r.systems}}<br>
                        <strong>Criterio de Éxito:</strong> ${{r.success}}
                    </div>
                </div>
            `).join('');
        }}

        // Populate Filtros
        function populateFilters() {{
            const selectCampo = document.getElementById('filterCampo');
            selectCampo.innerHTML = '<option value="">Todos los Campos de Acción</option>';
            camposData.forEach(c => {{
                let opt = document.createElement('option');
                opt.value = `CA-${{c.code.split('-')[1]}} ${{c.name}}`;
                opt.textContent = `${{c.code}} - ${{c.name}}`;
                selectCampo.appendChild(opt);
            }});

            const selectEslabon = document.getElementById('filterEslabon');
            selectEslabon.innerHTML = '<option value="">Todos los Eslabones</option>';
            eslabonesData.forEach(e => {{
                let opt = document.createElement('option');
                opt.value = e.name;
                opt.textContent = `${{e.code}} - ${{e.name}}`;
                selectEslabon.appendChild(opt);
            }});

            const selectEntidad = document.getElementById('filterEntidad');
            selectEntidad.innerHTML = '<option value="">Todas las Entidades Maestras</option>';
            let entidadesUnicas = [...new Set(attributesData.map(a => a.entidad).filter(e => e && e !== 'N/A'))].sort();
            entidadesUnicas.forEach(ent => {{
                let opt = document.createElement('option');
                opt.value = ent;
                opt.textContent = ent;
                selectEntidad.appendChild(opt);
            }});
        }}

        // Render Tabla Inventario
        function renderInventory(data) {{
            const tbody = document.getElementById('inventoryTableBody');
            tbody.innerHTML = '';

            if (data.length === 0) {{
                tbody.innerHTML = `<tr><td colspan="7" style="text-align:center; padding:2rem; color:var(--sda-text-muted);">No se encontraron atributos con los criterios seleccionados.</td></tr>`;
                return;
            }}

            data.forEach(item => {{
                let tr = document.createElement('tr');
                let badgeClass = 'badge-original';
                if (item.estado === 'DATO HOMOLOGADO') badgeClass = 'badge-homologado';
                if (item.estado === 'DATO VALIDADO') badgeClass = 'badge-validado';

                tr.innerHTML = `
                    <td style="font-weight:bold; color:var(--sda-dark-green);">${{item.id}}</td>
                    <td><strong>${{item.nombre}}</strong></td>
                    <td><span class="badge badge-entidad">${{item.entidad || 'N/A'}}</span></td>
                    <td>${{item.campo}}</td>
                    <td>${{item.eslabon}}</td>
                    <td>${{item.sistema}}</td>
                    <td><span class="badge ${{badgeClass}}">${{item.estado}}</span></td>
                `;
                tr.onclick = () => openModal(item);
                tbody.appendChild(tr);
            }});
        }}

        // Filtrar Inventario
        function filterInventory() {{
            const search = document.getElementById('searchInput').value.toLowerCase();
            const campo = document.getElementById('filterCampo').value;
            const eslabon = document.getElementById('filterEslabon').value;
            const entidad = document.getElementById('filterEntidad').value;
            const estado = document.getElementById('filterEstado').value;

            const filtered = attributesData.filter(item => {{
                const matchesSearch = !search || 
                    item.id.toLowerCase().includes(search) || 
                    item.nombre.toLowerCase().includes(search) || 
                    (item.entidad && item.entidad.toLowerCase().includes(search)) || 
                    item.fuente.toLowerCase().includes(search) || 
                    item.sistema.toLowerCase().includes(search);

                const matchesCampo = !campo || item.campo.includes(campo.split(' ')[0]);
                const matchesEslabon = !eslabon || item.eslabon.includes(eslabon.split(' ')[0]);
                const matchesEntidad = !entidad || item.entidad === entidad;
                const matchesEstado = !estado || item.estado === estado;

                return matchesSearch && matchesCampo && matchesEslabon && matchesEntidad && matchesEstado;
            }});

            renderInventory(filtered);
        }}

        // Open Modal Detalle
        function openModal(item) {{
            document.getElementById('modalTitle').textContent = `${{item.id}} - ${{item.nombre}}`;
            const modalBody = document.getElementById('modalBody');
            
            modalBody.innerHTML = `
                <div class="detail-grid">
                    <div class="detail-item"><div class="detail-label">ID Atributo</div><div class="detail-value">${{item.id}}</div></div>
                    <div class="detail-item"><div class="detail-label">Nombre / Atributo Base</div><div class="detail-value">${{item.nombre}}</div></div>
                    <div class="detail-item"><div class="detail-label">Entidad Maestra</div><div class="detail-value" style="color:var(--sda-dark-green); font-weight:bold;">${{item.entidad || 'N/A'}}</div></div>
                    <div class="detail-item"><div class="detail-label">Campo de Acción</div><div class="detail-value">${{item.campo}}</div></div>
                    <div class="detail-item"><div class="detail-label">Eslabón Cadena</div><div class="detail-value">${{item.eslabon}}</div></div>
                    <div class="detail-item"><div class="detail-label">Naturaleza del Dato</div><div class="detail-value">${{item.naturaleza}}</div></div>
                    <div class="detail-item"><div class="detail-label">Grano (Unidad)</div><div class="detail-value">${{item.grano}}</div></div>
                    <div class="detail-item"><div class="detail-label">Frecuencia</div><div class="detail-value">${{item.frecuencia}}</div></div>
                    <div class="detail-item"><div class="detail-label">Fuente Originaria</div><div class="detail-value">${{item.fuente}}</div></div>
                    <div class="detail-item"><div class="detail-label">Sistema Fuente</div><div class="detail-value">${{item.sistema}}</div></div>
                    <div class="detail-item"><div class="detail-label">Área Responsable</div><div class="detail-value">${{item.area}}</div></div>
                    <div class="detail-item"><div class="detail-label">Mecanismo Acceso</div><div class="detail-value">${{item.acceso}}</div></div>
                    <div class="detail-item"><div class="detail-label">Estado del Dato</div><div class="detail-value">${{item.estado}}</div></div>
                    <div class="detail-item"><div class="detail-label">Nivel Llave Analítica</div><div class="detail-value">${{item.llave}}</div></div>
                    <div class="detail-item"><div class="detail-label">¿Se Replica?</div><div class="detail-value">${{item.replica}}</div></div>
                    <div class="detail-item"><div class="detail-label">Grupo Acceso OTI</div><div class="detail-value">${{item.grupo}}</div></div>
                    <div class="detail-item"><div class="detail-label">Preguntas que Responde</div><div class="detail-value">${{item.preguntas}}</div></div>
                </div>
            `;
            document.getElementById('detailModal').style.display = 'flex';
        }}

        function closeModal() {{
            document.getElementById('detailModal').style.display = 'none';
        }}

        // Export Data
        function exportData(format) {{
            if (format === 'json') {{
                const dataStr = "data:text/json;charset=utf-8," + encodeURIComponent(JSON.stringify(attributesData, null, 2));
                const dlAnchorElem = document.createElement('a');
                dlAnchorElem.setAttribute("href", dataStr);
                dlAnchorElem.setAttribute("download", "Modelo_Metadatos_PGAE_SDA.json");
                dlAnchorElem.click();
            }} else if (format === 'csv') {{
                const headers = ["ID", "Nombre", "Entidad Maestra", "Campo", "Eslabón", "Naturaleza", "Grano", "Frecuencia", "Fuente", "Sistema", "Estado", "Responsable"];
                let csvContent = "data:text/csv;charset=utf-8," + headers.join(",") + "\\n";
                attributesData.forEach(row => {{
                    csvContent += `"${{row.id}}","${{row.nombre}}","${{row.entidad || ''}}","${{row.campo}}","${{row.eslabon}}","${{row.naturaleza}}","${{row.grano}}","${{row.frecuencia}}","${{row.fuente}}","${{row.sistema}}","${{row.estado}}","${{row.area}}"\\n`;
                }});
                const encodedUri = encodeURI(csvContent);
                const link = document.createElement("a");
                link.setAttribute("href", encodedUri);
                link.setAttribute("download", "Modelo_Metadatos_PGAE_SDA.csv");
                document.body.appendChild(link);
                link.click();
                document.body.removeChild(link);
            }}
        }}

        // Render Cards Grids
        function renderCampos() {{
            document.getElementById('camposGrid').innerHTML = camposData.map(c => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">${{c.icon || '📌'}}</div>
                        <div>
                            <div class="info-title">${{c.code}} - ${{c.name}}</div>
                            <div class="info-sub">Ámbito: ${{c.scope}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <p>${{c.desc}}</p>
                        <p style="margin-top:0.5rem; font-weight:bold; color:var(--sda-dark-green);">Total Atributos: ${{c.count}}</p>
                    </div>
                </div>
            `).join('');
        }}

        function renderEslabones() {{
            document.getElementById('eslabonesGrid').innerHTML = eslabonesData.map(e => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">${{e.icon || '🔹'}}</div>
                        <div>
                            <div class="info-title">${{e.code}} - ${{e.name}}</div>
                            <div class="info-sub">Total Atributos: ${{e.count}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <strong>Entregable del Eslabón:</strong>
                        <p>${{e.deliverable}}</p>
                    </div>
                </div>
            `).join('');
        }}

        function renderPreguntas() {{
            document.getElementById('preguntasGrid').innerHTML = preguntasData.map(p => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">❓</div>
                        <div>
                            <div class="info-title">${{p.code}} (${{p.topic}})</div>
                            <div class="info-sub">Campo Principal: ${{p.ca}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <p style="font-weight:bold; color:var(--sda-text-main);">${{p.question}}</p>
                    </div>
                </div>
            `).join('');
        }}

        function renderDimensiones() {{
            document.getElementById('dimensionesGrid').innerHTML = dimensionesData.map(d => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">🧩</div>
                        <div>
                            <div class="info-title">${{d.id}} - ${{d.name}}</div>
                            <div class="info-sub">Sistema Maestro: ${{d.system}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <p><strong>Concepto:</strong> ${{d.concept}}</p>
                        <p style="margin-top:0.4rem;"><strong>Atributos Clave:</strong> ${{d.attributes}}</p>
                        <p style="margin-top:0.4rem;"><strong>Área Propietaria:</strong> ${{d.owner}}</p>
                        <p style="margin-top:0.4rem; color:var(--sda-dark-green);"><strong>Usos:</strong> ${{d.uses}}</p>
                    </div>
                </div>
            `).join('');
        }}

        function renderIndicadores() {{
            document.getElementById('indicadoresGrid').innerHTML = indicadoresData.map(i => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">📈</div>
                        <div>
                            <div class="info-title">${{i.name}}</div>
                            <div class="info-sub">${{i.ca}} | Frecuencia: ${{i.freq}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <strong>Fórmula de Cálculo:</strong>
                        <div class="formula-box">${{i.formula}}</div>
                        <p><strong>Datos Requeridos:</strong> ${{i.inputs}}</p>
                        <p style="margin-top:0.4rem;"><strong>Responsable:</strong> ${{i.role}}</p>
                        <p style="margin-top:0.4rem; color:var(--sda-dark-green);"><strong>Uso Regulatorio:</strong> ${{i.use}}</p>
                    </div>
                </div>
            `).join('');
        }}

        function renderFuentes() {{
            document.getElementById('fuentesGrid').innerHTML = fuentesData.map(f => `
                <div class="info-card">
                    <div class="info-header">
                        <div class="info-icon">💻</div>
                        <div>
                            <div class="info-title">${{f.name}}</div>
                            <div class="info-sub">Tipo: ${{f.type}} | ${{f.attrs}}</div>
                        </div>
                    </div>
                    <div class="info-body">
                        <p><strong>Mecanismo Acceso:</strong> ${{f.mechanism}}</p>
                        <p style="margin-top:0.4rem;"><strong>Área Responsable:</strong> ${{f.owner}}</p>
                        <p style="margin-top:0.4rem;"><strong>Grupo Acceso OTI:</strong> ${{f.group}}</p>
                        <p style="margin-top:0.4rem;"><strong>Estado:</strong> <span class="badge badge-validado">${{f.status}}</span></p>
                        <p style="margin-top:0.4rem;"><strong>Responsable Gestión:</strong> ${{f.lead}}</p>
                    </div>
                </div>
            `).join('');
        }}

        // Gráficos Chart.js
        function initCharts() {{
            if (chartCamposInstance) chartCamposInstance.destroy();
            if (chartEslabonesInstance) chartEslabonesInstance.destroy();

            const ctxCampos = document.getElementById('chartCampos').getContext('2d');
            chartCamposInstance = new Chart(ctxCampos, {{
                type: 'bar',
                data: {{
                    labels: camposData.map(c => c.code),
                    datasets: [{{
                        label: 'Número de Atributos',
                        data: camposData.map(c => c.count),
                        backgroundColor: '#577515',
                        borderRadius: 6
                    }}]
                }},
                options: {{
                    responsive: true,
                    plugins: {{ legend: {{ display: false }} }},
                    scales: {{ y: {{ beginAtZero: true }} }}
                }}
            }});

            const ctxEslabones = document.getElementById('chartEslabones').getContext('2d');
            chartEslabonesInstance = new Chart(ctxEslabones, {{
                type: 'line',
                data: {{
                    labels: eslabonesData.map(e => e.code),
                    datasets: [{{
                        label: 'Atributos por Eslabón',
                        data: eslabonesData.map(e => e.count),
                        borderColor: '#89BC24',
                        backgroundColor: 'rgba(137, 188, 36, 0.2)',
                        fill: true,
                        tension: 0.3,
                        pointRadius: 5,
                        pointBackgroundColor: '#577515'
                    }}]
                }},
                options: {{
                    responsive: true,
                    plugins: {{ legend: {{ display: false }} }},
                    scales: {{ y: {{ beginAtZero: true }} }}
                }}
            }});
        }}

        // Handle Live Excel Upload directly in Browser via SheetJS
        function handleExcelUpload(event) {{
            const file = event.target.files[0];
            if (!file) return;

            if (typeof XLSX === 'undefined') {{
                alert("⚠️ La librería SheetJS no está cargada. Asegúrate de tener conexión a Internet o usa 'python generar_dashboard.py' en la consola.");
                event.target.value = '';
                return;
            }}

            const reader = new FileReader();
            reader.onload = function(e) {{
                try {{
                    const data = new Uint8Array(e.target.result);
                    const workbook = XLSX.read(data, {{ type: 'array' }});
                    
                    console.log("Hojas detectadas:", workbook.SheetNames);

                    // 1. Parse Inventory Sheet
                    let invSheetName = workbook.SheetNames.find(s => s.toLowerCase().includes('inventario')) || 
                                       workbook.SheetNames.find(s => s.toLowerCase().includes('pgae')) || 
                                       workbook.SheetNames[0];
                    let wsInv = workbook.Sheets[invSheetName];
                    let invRows = XLSX.utils.sheet_to_json(wsInv, {{ header: 1 }});

                    // Detect Dynamic Header Map
                    let colMap = {{}};
                    let headerRow = invRows.find(r => r && r.some(c => String(c).toLowerCase().includes('id atributo') || String(c).toLowerCase().includes('eslabón') || String(c).toLowerCase().includes('eslabon')));

                    if (headerRow) {{
                        headerRow.forEach((val, idx) => {{
                            if (!val) return;
                            let v = String(val).toLowerCase();
                            if (v.includes('id') && v.includes('atributo')) colMap.id = idx;
                            else if (v.includes('eslab')) colMap.eslabon = idx;
                            else if (v.includes('campo')) colMap.campo = idx;
                            else if (v.includes('entidad')) colMap.entidad = idx;
                            else if (v.includes('dato') || v.includes('atributo base')) colMap.nombre = idx;
                            else if (v.includes('naturaleza')) colMap.naturaleza = idx;
                            else if (v.includes('grano')) colMap.grano = idx;
                            else if (v.includes('frecuencia')) colMap.frecuencia = idx;
                            else if (v.includes('fuente') && v.includes('origin')) colMap.fuente = idx;
                            else if (v.includes('sistema')) colMap.sistema = idx;
                            else if (v.includes('área') || v.includes('area')) colMap.area = idx;
                            else if (v.includes('mecanismo')) colMap.acceso = idx;
                            else if (v.includes('estado')) colMap.estado = idx;
                            else if (v.includes('llave') || v.includes('nivel')) colMap.llave = idx;
                            else if (v.includes('replica')) colMap.replica = idx;
                            else if (v.includes('grupo')) colMap.grupo = idx;
                            else if (v.includes('pregunta')) colMap.preguntas = idx;
                        }});
                    }}

                    function getCellVal(row, key, defaultIdx) {{
                        let idx = (colMap[key] !== undefined) ? colMap[key] : defaultIdx;
                        return String(row[idx] || '').trim();
                    }}

                    let hasEntidadCol = colMap.entidad !== undefined;
                    let offset = hasEntidadCol ? 1 : 0;

                    let newAttrs = [];
                    for (let r = 0; r < invRows.length; r++) {{
                        let row = invRows[r];
                        let idIdx = colMap.id !== undefined ? colMap.id : 0;
                        if (row && row[idIdx] && String(row[idIdx]).trim().toUpperCase().startsWith('ATR-')) {{
                            newAttrs.push({{
                                id: String(row[idIdx]).trim(),
                                eslabon: getCellVal(row, 'eslabon', 1),
                                campo: getCellVal(row, 'campo', 2),
                                entidad: hasEntidadCol ? getCellVal(row, 'entidad', 3) : 'N/A',
                                nombre: getCellVal(row, 'nombre', 3 + offset),
                                naturaleza: getCellVal(row, 'naturaleza', 4 + offset),
                                grano: getCellVal(row, 'grano', 5 + offset),
                                frecuencia: getCellVal(row, 'frecuencia', 6 + offset),
                                fuente: getCellVal(row, 'fuente', 7 + offset),
                                sistema: getCellVal(row, 'sistema', 8 + offset),
                                area: getCellVal(row, 'area', 9 + offset),
                                acceso: getCellVal(row, 'acceso', 10 + offset),
                                estado: getCellVal(row, 'estado', 11 + offset),
                                llave: getCellVal(row, 'llave', 12 + offset),
                                replica: getCellVal(row, 'replica', 13 + offset),
                                grupo: getCellVal(row, 'grupo', 14 + offset),
                                preguntas: getCellVal(row, 'preguntas', 15 + offset)
                            }});
                        }}
                    }}

                    if (newAttrs.length === 0) {{
                        alert("⚠️ No se encontraron atributos en formato 'ATR-...' en la hoja " + invSheetName + ".");
                        event.target.value = '';
                        return;
                    }}

                    attributesData = newAttrs;

                    // 2. Parse Matrix Sheet if available
                    let matSheetName = workbook.SheetNames.find(s => s.toLowerCase().includes('matriz'));
                    const campos = ['CA-01', 'CA-02', 'CA-03', 'CA-04', 'CA-05', 'CA-06', 'CA-07', 'CA-08', 'CA-09'];
                    const eslabones = ['E1', 'E2', 'E3', 'E4', 'E5', 'E6', 'E7', 'E8'];

                    let newMatrix = {{}};
                    if (matSheetName) {{
                        let wsMat = workbook.Sheets[matSheetName];
                        let matRows = XLSX.utils.sheet_to_json(wsMat, {{ header: 1 }});
                        
                        eslabones.forEach(e => {{
                            newMatrix[e] = {{}};
                            let rowMatch = matRows.find(r => r && r[0] && String(r[0]).trim().startsWith(e));
                            campos.forEach((c, idx) => {{
                                let val = rowMatch ? Number(rowMatch[idx + 1] || 0) : 0;
                                newMatrix[e][c] = isNaN(val) ? 0 : val;
                            }});
                        }});
                    }} else {{
                        eslabones.forEach(e => {{
                            newMatrix[e] = {{}};
                            campos.forEach(c => {{
                                newMatrix[e][c] = attributesData.filter(a => 
                                    a.eslabon.toUpperCase().includes(e) && a.campo.toUpperCase().includes(c)
                                ).length;
                            }});
                        }});
                    }}
                    matrixData = newMatrix;

                    // 3. Update Counts
                    camposData.forEach(c => {{
                        c.count = attributesData.filter(a => a.campo.toUpperCase().includes(c.code)).length;
                    }});

                    eslabonesData.forEach(e => {{
                        e.count = attributesData.filter(a => a.eslabon.toUpperCase().includes(e.code)).length;
                    }});

                    // 4. Parse Optional Sheets
                    let pregSheet = workbook.SheetNames.find(s => s.toLowerCase().includes('preguntas'));
                    if (pregSheet) {{
                        let pRows = XLSX.utils.sheet_to_json(workbook.Sheets[pregSheet], {{ header: 1 }});
                        let newPreg = [];
                        for (let r = 0; r < pRows.length; r++) {{
                            let row = pRows[r];
                            if (row && row[0] && String(row[0]).trim().startsWith('PE-')) {{
                                newPreg.push({{
                                    code: String(row[0]).trim(),
                                    topic: String(row[1] || '').trim(),
                                    question: String(row[2] || '').trim(),
                                    ca: String(row[3] || '').trim()
                                }});
                            }}
                        }}
                        if (newPreg.length > 0) preguntasData = newPreg;
                    }}

                    let dimSheet = workbook.SheetNames.find(s => s.toLowerCase().includes('dimensiones'));
                    if (dimSheet) {{
                        let dRows = XLSX.utils.sheet_to_json(workbook.Sheets[dimSheet], {{ header: 1 }});
                        let newDim = [];
                        for (let r = 0; r < dRows.length; r++) {{
                            let row = dRows[r];
                            if (row && row[0] && String(row[0]).trim().startsWith('DM-')) {{
                                newDim.push({{
                                    id: String(row[0]).trim(),
                                    name: String(row[1] || '').trim(),
                                    concept: String(row[2] || '').trim(),
                                    attributes: String(row[3] || '').trim(),
                                    system: String(row[4] || '').trim(),
                                    owner: String(row[5] || '').trim(),
                                    uses: String(row[6] || '').trim()
                                }});
                            }}
                        }}
                        if (newDim.length > 0) dimensionesData = newDim;
                    }}

                    let indSheet = workbook.SheetNames.find(s => s.toLowerCase().includes('indicadores'));
                    if (indSheet) {{
                        let iRows = XLSX.utils.sheet_to_json(workbook.Sheets[indSheet], {{ header: 1 }});
                        let newInd = [];
                        for (let r = 3; r < iRows.length; r++) {{
                            let row = iRows[r];
                            if (row && row[0]) {{
                                newInd.push({{
                                    name: String(row[0]).trim(),
                                    ca: String(row[1] || '').trim(),
                                    formula: String(row[2] || '').trim(),
                                    inputs: String(row[3] || '').trim(),
                                    freq: String(row[4] || '').trim(),
                                    role: String(row[5] || '').trim(),
                                    use: String(row[6] || '').trim()
                                }});
                            }}
                        }}
                        if (newInd.length > 0) indicadoresData = newInd;
                    }}

                    let fueSheet = workbook.SheetNames.find(s => s.toLowerCase().includes('fuentes'));
                    if (fueSheet) {{
                        let fRows = XLSX.utils.sheet_to_json(workbook.Sheets[fueSheet], {{ header: 1 }});
                        let newFue = [];
                        for (let r = 3; r < fRows.length; r++) {{
                            let row = fRows[r];
                            if (row && row[0] != null) {{
                                newFue.push({{
                                    id: String(row[0]).trim(),
                                    name: String(row[1] || '').trim(),
                                    type: String(row[2] || '').trim(),
                                    attrs: String(row[3] || '').trim(),
                                    owner: String(row[4] || '').trim(),
                                    mechanism: String(row[5] || '').trim(),
                                    group: String(row[6] || '').trim(),
                                    status: String(row[7] || '').trim(),
                                    lead: String(row[8] || '').trim()
                                }});
                            }}
                        }}
                        if (newFue.length > 0) fuentesData = newFue;
                    }}

                    // Refrescar la interfaz completa
                    refreshAllUI();
                    
                    alert(`✅ ¡Excelente! Se actualizó correctamente el dashboard con ${{newAttrs.length}} atributos del archivo "${{file.name}}".`);
                }} catch (err) {{
                    console.error("Error al procesar Excel:", err);
                    alert("❌ Ocurrió un error al leer el archivo Excel: " + err.message);
                }} finally {{
                    event.target.value = '';
                }}
            }};
            reader.readAsArrayBuffer(file);
        }}
    </script>
</body>
</html>
"""

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print(f"OK -> Dashboard completado e index.html escrito con exito en {output_path}!")

def main():
    workspace_dir = r"h:\Mi unidad\SECRETARIA DISTRITAL AMBIENTE\2026\MODELO DE DATOS PGAE"
    
    # Buscar el archivo Excel de metadatos más reciente en la carpeta
    excel_candidates = sorted(glob.glob(os.path.join(workspace_dir, "Modelo_Metadatos_PGAE_SDA*.xlsx")), key=os.path.getmtime, reverse=True)
    excel_candidates = [f for f in excel_candidates if '~$' not in f]
    
    default_excel = excel_candidates[0] if excel_candidates else os.path.join(workspace_dir, "Modelo_Metadatos_PGAE_SDA-v10.xlsx")
    
    excel_path = sys.argv[1] if len(sys.argv) > 1 else default_excel
    if not os.path.exists(excel_path):
        print(f"ERROR: El archivo Excel no existe en: {excel_path}")
        sys.exit(1)
        
    output_html = os.path.join(workspace_dir, "index.html")
    
    data = parse_excel_metadata(excel_path)
    generate_html_dashboard(data, output_html)

if __name__ == "__main__":
    main()
